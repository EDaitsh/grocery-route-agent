"""הרחבות האייג'נט (לא LLM):

1. StoreRAG  - שליפה (Retrieval) מהחנות: מחלקות (עם דוגמאות למה שנמכר בהן) ורשימה אופציונלית של
               מוצרים ידועים - בעיקר חריגים שלא נמצאים במחלקה הצפויה. אינדקס TF-IDF על n-grams של תווים,
               שעמיד לשגיאות כתיב ולתחיליות בעברית, ופועל מקומית בלי API של embeddings.
2. StoreMap  - כלי פנימי (Route Tool): חישוב מסלול קצר בחנות. Dijkstra בין נקודות + nearest-neighbor
               ו-2-opt לסדר העצירות, עם אילוץ: מוצרים שמסומנים "לאסוף בסוף" (גלידה) אחרונים לפני הקופות.
               המפה נטענת מציור בתאי אקסל (store.xlsx) או מגרף ב-map.json.
3. save_text - כתיבת קובץ (פעולה רגישה, מופעלת רק אחרי אישור HITL בגרף).
"""
from __future__ import annotations

import csv
import difflib
import heapq
import json
import math
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from . import config

# ---------------------------------------------------------------- נרמול טקסט עברי
_FINALS = str.maketrans({"ך": "כ", "ם": "מ", "ן": "נ", "ף": "פ", "ץ": "צ"})
_NIQQUD = re.compile(r"[֑-ׇ]")
_PUNCT = re.compile(r"[\"'׳״`.,;:!?()\[\]{}\-_/\\]+")


def normalize(text: str) -> str:
    text = _NIQQUD.sub("", text.lower())
    text = _PUNCT.sub(" ", text).translate(_FINALS)
    return re.sub(r"\s+", " ", text).strip()


def _ngrams(text: str, ns=(2, 3)) -> Counter:
    grams: Counter = Counter()
    for word in normalize(text).split():
        padded = f" {word} "
        for n in ns:
            for i in range(len(padded) - n + 1):
                grams[padded[i:i + n]] += 1
    return grams


# ---------------------------------------------------------------- RAG
@dataclass
class Product:
    id: str
    name: str
    aliases: list[str]
    category: str
    location: str
    collect_last: bool   # מוצר שנמס (גלידה): האזור שלו נדחה לסוף המסלול

    def as_dict(self) -> dict:
        return {"id": self.id, "name": self.name, "category": self.category,
                "location": self.location, "collect_last": self.collect_last}


@dataclass
class Department:
    name: str
    location: str
    examples: list[str]   # מה נמכר במחלקה. משמש את ה-LLM לשיוך ואת ה-RAG לשליפה

    def as_dict(self) -> dict:
        return {"name": self.name, "location": self.location}


def read_catalog_csv(path: Path) -> list[Product]:
    with open(path, encoding="utf-8") as f:
        return [Product(row["id"], row["name"], [a for a in row["aliases"].split("|") if a],
                        row["category"], row["location"], row.get("collect_last") == "1")
                for row in csv.DictReader(f)]


def departments_from_products(products: list[Product]) -> list[Department]:
    """לחנות בפורמט הישן (קטלוג מלא בלי מחלקות): כל קטגוריה היא מחלקה, במיקום הנפוץ של המוצרים שלה."""
    by_cat: dict[str, list[Product]] = {}
    for p in products:
        by_cat.setdefault(p.category, []).append(p)
    return [Department(cat, Counter(p.location for p in ps).most_common(1)[0][0],
                       list(dict.fromkeys(v for p in ps for v in (p.name, *p.aliases))))
            for cat, ps in by_cat.items()]


class _CharIndex:
    """TF-IDF על n-grams של תווים. כל וריאציית שם היא "מסמך" נפרד, והציון של רשומה הוא המקסימום ביניהם."""

    def __init__(self, entries: dict[str, list[str]]):
        docs: list[tuple[str, Counter]] = []
        df: Counter = Counter()
        self._exact: dict[str, list[str]] = {}
        for key, variants in entries.items():
            for variant in variants:
                grams = _ngrams(variant)
                docs.append((key, grams))
                df.update(grams.keys())
                self._exact.setdefault(normalize(variant), []).append(key)
        self._idf = {g: math.log((1 + len(docs)) / (1 + c)) + 1 for g, c in df.items()}
        self._vecs = [(key, self._weigh(g)) for key, g in docs]

    def _weigh(self, grams: Counter) -> dict[str, float]:
        vec = {g: c * self._idf.get(g, 1.0) for g, c in grams.items()}
        norm = math.sqrt(sum(v * v for v in vec.values())) or 1.0
        return {g: v / norm for g, v in vec.items()}

    def search(self, query: str, k: int) -> list[tuple[str, float]]:
        q = self._weigh(_ngrams(query))
        best: dict[str, float] = {}
        for key, vec in self._vecs:
            s = sum(w * vec.get(g, 0.0) for g, w in q.items())
            if s > best.get(key, 0.0):
                best[key] = s
        # התאמה מדויקת לשם או לכינוי מקבלת ציון מלא
        for key in dict.fromkeys(self._exact.get(normalize(query), [])):
            best[key] = 1.0
        ranked = sorted(best.items(), key=lambda x: -x[1])[:k]
        return [(key, round(score, 3)) for key, score in ranked if score > 0.05]


class StoreRAG:
    def __init__(self, departments: list[Department], products: list[Product] = ()):
        self.departments: dict[str, Department] = {d.name: d for d in departments}
        self.products: dict[str, Product] = {p.id: p for p in products}
        self._dep_index = _CharIndex({d.name: [d.name, *d.examples] for d in departments})
        self._prod_index = _CharIndex({p.id: [p.name, *p.aliases] for p in products})

    def search(self, query: str, k: int = 3) -> list[dict]:
        """מחזיר עד k מוצרים ידועים מועמדים עם ציון 0..1."""
        return [{**self.products[pid].as_dict(), "score": s} for pid, s in self._prod_index.search(query, k)]

    def search_departments(self, query: str, k: int = 3) -> list[dict]:
        """מחזיר עד k מחלקות שהדוגמאות שלהן הכי דומות לפריט, עם ציון 0..1."""
        return [{**self.departments[n].as_dict(), "score": s} for n, s in self._dep_index.search(query, k)]

    def get(self, pid: str) -> Product | None:
        return self.products.get(pid)

    def department(self, name: str) -> Department | None:
        return self.departments.get(name)


# ---------------------------------------------------------------- מפה ומסלול
@dataclass
class RouteResult:
    stops: list[dict] = field(default_factory=list)   # [{node, label, products:[...], walk_m}]
    total_m: float = 0.0
    naive_m: float = 0.0          # מרחק אם הולכים לפי סדר הרשימה המקורי
    path_nodes: list[str] = field(default_factory=list)


class StoreMap:
    """גרף הליכה בחנות. כל אזור (עצירה אפשרית) ממופה לצומת בגרף."""

    def __init__(self, name: str, adj: dict[str, list[tuple[str, float]]], labels: dict[str, str],
                 start: str, end: str, zone_node: dict[str, str] | None = None):
        self.name = name
        self.adj = adj
        self.labels = labels                          # אזור -> שם לתצוגה
        self.start, self.end = start, end
        self._zone_node = zone_node or {z: z for z in labels}
        self._sp_cache: dict[str, tuple[dict, dict]] = {}

    @classmethod
    def from_json(cls, map_path: Path) -> StoreMap:
        data = json.loads(Path(map_path).read_text(encoding="utf-8"))
        nodes = data["nodes"]
        adj: dict[str, list[tuple[str, float]]] = {n: [] for n in nodes}
        for a, b, *w in data["edges"]:
            d = w[0] if w else math.hypot(nodes[a]["x"] - nodes[b]["x"], nodes[a]["y"] - nodes[b]["y"])
            adj[a].append((b, d))
            adj[b].append((a, d))
        return cls(data["store_name"], adj, {n: v["label"] for n, v in nodes.items()}, data["start"], data["end"])

    @property
    def zones(self) -> set[str]:
        return set(self.labels)

    def label(self, zone: str) -> str:
        return self.labels[zone]

    def _dijkstra(self, src: str) -> tuple[dict, dict]:
        if src in self._sp_cache:
            return self._sp_cache[src]
        dist, prev = {src: 0.0}, {}
        pq = [(0.0, src)]
        while pq:
            d, u = heapq.heappop(pq)
            if d > dist.get(u, math.inf):
                continue
            for v, w in self.adj[u]:
                nd = d + w
                if nd < dist.get(v, math.inf):
                    dist[v], prev[v] = nd, u
                    heapq.heappush(pq, (nd, v))
        self._sp_cache[src] = (dist, prev)
        return dist, prev

    def dist(self, a: str, b: str) -> float:
        return self._dijkstra(self._zone_node[a])[0].get(self._zone_node[b], math.inf)

    def path(self, a: str, b: str) -> list[str]:
        a, b = self._zone_node[a], self._zone_node[b]
        _, prev = self._dijkstra(a)
        out = [b]
        while out[-1] != a:
            out.append(prev[out[-1]])
        return out[::-1]

    def _tour_len(self, seq: list[str]) -> float:
        return sum(self.dist(seq[i], seq[i + 1]) for i in range(len(seq) - 1))

    def _order(self, stops: list[str], start: str, next_fixed: str) -> list[str]:
        """nearest-neighbor ואחריו 2-opt, כשנקודת ההתחלה והנקודה שאחרי הקבוצה קבועות."""
        remaining, cur, order = set(stops), start, []
        while remaining:
            nxt = min(remaining, key=lambda s: (self.dist(cur, s), s))
            order.append(nxt)
            remaining.remove(nxt)
            cur = nxt
        improved = True
        while improved and len(order) > 2:
            improved = False
            for i in range(len(order) - 1):
                for j in range(i + 1, len(order)):
                    cand = order[:i] + order[i:j + 1][::-1] + order[j + 1:]
                    if self._tour_len([start, *cand, next_fixed]) + 1e-9 < self._tour_len([start, *order, next_fixed]):
                        order, improved = cand, True
        return order

    def plan(self, items: list[dict]) -> RouteResult:
        """items: [{"name", "location", "collect_last", "qty"}] -> מסלול מסודר."""
        by_node: dict[str, list[dict]] = {}
        for it in items:
            by_node.setdefault(it["location"], []).append(it)
        # אזור עם מוצר שנמס עובר לסוף, וכל המוצרים באותו אזור נאספים יחד איתו
        last = [n for n, its in by_node.items() if n != self.end and any(it.get("collect_last") for it in its)]
        regular = [n for n in by_node if n not in last and n != self.end]

        order = self._order(regular, self.start, last[0] if len(last) == 1 else self.end)
        order += self._order(last, order[-1] if order else self.start, self.end)
        full = [self.start, *order, self.end]

        res = RouteResult(total_m=round(self._tour_len(full), 1))
        for i in range(1, len(full)):
            seg = self.path(full[i - 1], full[i])
            res.path_nodes.extend(seg if i == 1 else seg[1:])
            res.stops.append({"node": full[i], "label": self.label(full[i]),
                              "products": by_node.get(full[i], []), "collect_last": full[i] in last,
                              "walk_m": round(self.dist(full[i - 1], full[i]), 1)})

        # להשוואה: הליכה לפי הסדר שבו המשתמש כתב את הרשימה
        naive, seen = [self.start], set()
        for it in items:
            if it["location"] not in seen:
                naive.append(it["location"])
                seen.add(it["location"])
        res.naive_m = round(self._tour_len([*naive, self.end]), 1)
        return res


# ---------------------------------------------------------------- קובץ חנות באקסל
# לשונית "מפת החנות": ציור בתאים, כל תא בערך מטר על מטר. תא ריק = מעבר הליכה, X = מדף/קיר,
# טקסט = אזור (המקום שבו הלקוח עומד כשהוא לוקח מוצרים). שורה 1 להגדרות, הציור מתחיל בשורה 3.
# לשונית "מחלקות": שורה לכל מחלקה, עמודת "אזור" מפנה לטקסט שבציור, ודוגמאות למה שנמכר בה.
# לשונית "חריגים" (לא חובה): מוצרים שלא נמצאים במחלקה הצפויה (למשל סוללות ליד הקופות).
# בפורמט הקודם הייתה לשונית "מוצרים" עם קטלוג מלא - היא עדיין נתמכת, והמחלקות נגזרות ממנה.
MAP_SHEET, DEPARTMENTS_SHEET, PRODUCTS_SHEET, LEGACY_PRODUCTS_SHEET = "מפת החנות", "מחלקות", "חריגים", "מוצרים"
_WALL = {"x", "×", "#"}
_ENTRANCE, _CHECKOUT = ("כניסה",), ("קופות", "קופה")
_YES = {"1", "כן", "v", "✓", "x", "true"}
_LIST_SEP = re.compile(r"\s*[,|،]\s*")
_PRODUCT_COLS = {"id": ("מק\"ט", "מק״ט", "מקט"), "name": ("שם המוצר",), "aliases": ("שמות נוספים",),
                 "category": ("קטגוריה",), "location": ("אזור",), "last": ("לאסוף בסוף",)}
_DEPARTMENT_COLS = {"name": ("מחלקה",), "location": ("אזור",), "examples": ("מה נמכר כאן", "דוגמאות")}


class StoreFileError(ValueError):
    def __init__(self, path: Path, problems: list[str]):
        self.problems = problems
        super().__init__(f"נמצאו בעיות בקובץ החנות {Path(path).name}:\n- " + "\n- ".join(problems))


def _clean(v) -> str:
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


def _grid_map(ws, default_name: str, problems: list[str]) -> StoreMap | None:
    name, cell_m = default_name, 1.0
    row1 = [c.value for c in ws[1]]
    for i, v in enumerate(row1):
        nxt = next((x for x in row1[i + 1:] if _clean(x)), None)
        if "שם החנות" in _clean(v) and nxt:
            name = _clean(nxt)
        elif "מטר" in _clean(v) and isinstance(nxt, (int, float)) and nxt > 0:
            cell_m = float(nxt)

    values: dict[tuple[int, int], str] = {}
    for row in ws.iter_rows(min_row=3):
        for c in row:
            if _clean(c.value):
                values[(c.row, c.column)] = _clean(c.value)
    for rng in ws.merged_cells.ranges:          # בתא ממוזג הערך נמצא רק בפינה השמאלית העליונה
        v = values.get((rng.min_row, rng.min_col))
        if v and rng.min_row >= 3:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    values[(r, c)] = v
    if not values:
        problems.append(f"הלשונית '{MAP_SHEET}' ריקה: לא נמצא ציור של החנות")
        return None

    rows, cols = [r for r, _ in values], [c for _, c in values]
    walk = {(r, c) for r in range(min(rows), max(rows) + 1) for c in range(min(cols), max(cols) + 1)
            if values.get((r, c), "").lower() not in _WALL}
    zone_cells: dict[str, list[tuple[int, int]]] = {}
    for cell in walk:
        if cell in values:
            zone_cells.setdefault(values[cell], []).append(cell)

    def node(cell):
        return f"{cell[0]},{cell[1]}"

    adj: dict[str, list[tuple[str, float]]] = {node(c): [] for c in walk}
    for r, c in walk:
        for dr, dc in ((-1, -1), (-1, 0), (-1, 1), (0, -1), (0, 1), (1, -1), (1, 0), (1, 1)):
            nb = (r + dr, c + dc)
            # באלכסון רק אם שני התאים הצמודים פנויים, כדי לא לחתוך פינה של מדף
            if nb in walk and (not (dr and dc) or ((r + dr, c) in walk and (r, c + dc) in walk)):
                adj[node((r, c))].append((node(nb), cell_m * (math.sqrt(2) if dr and dc else 1)))

    def center(cells):
        cr, cc = sum(r for r, _ in cells) / len(cells), sum(c for _, c in cells) / len(cells)
        return node(min(cells, key=lambda x: ((x[0] - cr) ** 2 + (x[1] - cc) ** 2, x)))

    start = next((z for z in _ENTRANCE if z in zone_cells), None)
    end = next((z for z in _CHECKOUT if z in zone_cells), None)
    if not start:
        problems.append("לא נמצאה כניסה בציור: כתבו 'כניסה' בתא שבו נכנסים לחנות")
    if not end:
        problems.append("לא נמצאו קופות בציור: כתבו 'קופות' בתא שבו משלמים")
    if not (start and end):
        return None

    smap = StoreMap(name, adj, {z: z for z in zone_cells}, start, end,
                    {z: center(cells) for z, cells in zone_cells.items()})
    for z in sorted(zone_cells):
        if smap.dist(start, z) == math.inf:
            problems.append(f"אין דרך מהכניסה לאזור '{z}': בדקו שתאי X לא חוסמים אותו מכל הכיוונים")
    return smap


def _rows(ws, cols: dict[str, tuple[str, ...]], required: tuple[str, ...], problems: list[str]):
    """מחזיר (מספר שורה, get) לכל שורה לא ריקה. get(key) = ערך העמודה, או "" אם אין."""
    header = [_clean(c.value) for c in ws[1]]
    idx = {key: next((i for i, h in enumerate(header) if h in names), None) for key, names in cols.items()}
    missing = [key for key in required if idx[key] is None]
    for key in missing:
        problems.append(f"בלשונית '{ws.title}' חסרה עמודה בשם '{cols[key][0]}'")
    if missing:
        return
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not any(_clean(v) for v in row):
            continue
        # שורות "דוגמה: ..." בתבנית הן הסבר, לא נתונים
        if _clean(row[idx[required[0]]] if idx[required[0]] < len(row) else "").startswith("דוגמה"):
            continue
        yield n, lambda key, row=row: _clean(row[idx[key]]) if idx[key] is not None and idx[key] < len(row) else ""


def _check_zone(what: str, loc: str, smap: StoreMap | None, problems: list[str]) -> bool:
    if not loc:
        problems.append(f"{what}: לא נבחר אזור")
        return False
    if smap and loc not in smap.zones:
        hint = difflib.get_close_matches(loc, sorted(smap.zones), n=1)
        problems.append(f"{what} משויך לאזור '{loc}' שלא קיים בציור"
                        + (f" (אולי התכוונתם ל'{hint[0]}'?)" if hint else ""))
        return False
    return True


def _xlsx_departments(ws, smap: StoreMap | None, problems: list[str]) -> list[Department]:
    departments, seen = [], set()
    for n, get in _rows(ws, _DEPARTMENT_COLS, ("name", "location"), problems):
        name, loc = get("name") or get("location"), get("location")
        if not name:
            problems.append(f"מחלקות, שורה {n}: חסר שם מחלקה")
            continue
        if not _check_zone(f"המחלקה '{name}'", loc, smap, problems):
            continue
        if name in seen:
            problems.append(f"המחלקה '{name}' מופיעה יותר מפעם אחת")
            continue
        seen.add(name)
        departments.append(Department(name, loc, [e for e in _LIST_SEP.split(get("examples")) if e]))
    if not departments and not problems:
        problems.append(f"הלשונית '{ws.title}' ריקה: הגדירו לפחות מחלקה אחת")
    return departments


def _xlsx_products(ws, smap: StoreMap | None, problems: list[str]) -> list[Product]:
    products, seen = [], set()
    for n, get in _rows(ws, _PRODUCT_COLS, ("name", "location"), problems):
        name, loc = get("name"), get("location")
        if not name:
            problems.append(f"{ws.title}, שורה {n}: חסר שם מוצר")
            continue
        if not _check_zone(f"המוצר '{name}'", loc, smap, problems):
            continue
        pid = get("id") or f"R{n}"
        if pid in seen:
            problems.append(f"המק\"ט {pid} מופיע יותר מפעם אחת")
            continue
        seen.add(pid)
        products.append(Product(pid, name, [a.strip() for a in get("aliases").split("|") if a.strip()],
                                get("category") or loc, loc, get("last").lower() in _YES))
    return products


def load_store_xlsx(path: Path) -> tuple[StoreRAG, StoreMap]:
    """טוען חנות מקובץ אקסל. זורק StoreFileError עם כל הבעיות שנמצאו, בעברית פשוטה."""
    from openpyxl import load_workbook

    path = Path(path)
    wb = load_workbook(path, data_only=True)
    products_sheet = next((s for s in (PRODUCTS_SHEET, LEGACY_PRODUCTS_SHEET) if s in wb.sheetnames), None)
    missing = [s for s in (MAP_SHEET,) if s not in wb.sheetnames]
    if DEPARTMENTS_SHEET not in wb.sheetnames and products_sheet != LEGACY_PRODUCTS_SHEET:
        missing.append(DEPARTMENTS_SHEET)
    if missing:
        raise StoreFileError(path, [f"חסרה לשונית בשם '{s}'" for s in missing])
    problems: list[str] = []
    smap = _grid_map(wb[MAP_SHEET], path.parent.name, problems)
    products = _xlsx_products(wb[products_sheet], smap, problems) if products_sheet else []
    if DEPARTMENTS_SHEET in wb.sheetnames:
        departments = _xlsx_departments(wb[DEPARTMENTS_SHEET], smap, problems)
    else:
        departments = departments_from_products(products)
        if not departments and not problems:
            problems.append(f"הלשונית '{products_sheet}' ריקה")
    if problems:
        raise StoreFileError(path, problems)
    return StoreRAG(departments, products), smap


# ---------------------------------------------------------------- טעינת חנות
STORE_XLSX = "store.xlsx"
_STORE_CACHE: dict[str, tuple[StoreRAG, StoreMap]] = {}


def list_stores() -> list[str]:
    return sorted(p.name for p in config.STORES_DIR.iterdir()
                  if (p / STORE_XLSX).exists() or (p / "map.json").exists())


def load_store(name: str) -> tuple[StoreRAG, StoreMap]:
    if name not in _STORE_CACHE:
        d = config.STORES_DIR / name
        if (d / STORE_XLSX).exists():
            _STORE_CACHE[name] = load_store_xlsx(d / STORE_XLSX)
        elif (d / "map.json").exists():
            products = read_catalog_csv(d / "catalog.csv")
            _STORE_CACHE[name] = (StoreRAG(departments_from_products(products), products),
                                  StoreMap.from_json(d / "map.json"))
        else:
            raise FileNotFoundError(f"החנות '{name}' לא נמצאה. חנויות זמינות: {', '.join(list_stores())}")
    return _STORE_CACHE[name]


# ---------------------------------------------------------------- קבצים
def read_list_file(path: str | Path) -> str:
    p = Path(str(path).strip().strip('"\'')).expanduser()
    if not p.exists():
        raise FileNotFoundError(f"הקובץ לא נמצא: {p}")
    if p.suffix.lower() not in {".txt", ".csv", ".md", ""}:
        raise ValueError("ניתן לטעון רק קבצי טקסט (txt / csv / md).")
    if p.stat().st_size > 100_000:
        raise ValueError("הקובץ גדול מדי לרשימת קניות (מעל 100KB).")
    return p.read_text(encoding="utf-8")


def save_text(path: str | Path, text: str) -> Path:
    # Users often paste paths wrapped in quotes (Windows "Copy as path")
    p = Path(str(path).strip().strip('"\'')).expanduser()
    if p.is_dir():
        p = p / f"route_{datetime.now():%Y%m%d_%H%M}.txt"
    elif not p.suffix:
        p = p.with_suffix(".txt")
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(text, encoding="utf-8")
    return p.resolve()
